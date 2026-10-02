"""API tests: run lifecycle, review routing, reviewer actions, and end-to-end PII absence
in API responses, the database, logs and exports."""
from __future__ import annotations

import copy
import dataclasses
import io
import logging
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import inspect, text

from app.db.base import get_engine
from app.main import create_app


class ListHandler(logging.Handler):
    def __init__(self):
        super().__init__(logging.DEBUG)
        self.lines: list[str] = []

    def emit(self, record):
        try:
            self.lines.append(self.format(record))
        except Exception:
            self.lines.append(str(record.msg))


@pytest.fixture(scope="module")
def env(settings, fixture_bundle, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("api")
    app_cfg = copy.deepcopy(settings.app)
    app_cfg["data_source"].update(fixtures_path=str(fixture_bundle[0]), uploads_dir=str(tmp / "uploads"),
                                  output_dir=str(tmp / "output"))
    s = dataclasses.replace(settings, app=app_cfg, database_url=f"sqlite:///{(tmp / 'test.db').as_posix()}")
    handler = ListHandler()
    root = logging.getLogger()
    root.addHandler(handler)
    prev = root.level
    root.setLevel(logging.DEBUG)
    with TestClient(create_app(s)) as client:
        r = client.post("/api/runs", json={"source": "fixtures"})
        assert r.status_code == 202
        run_id = r.json()["id"]
        deadline = time.time() + 120
        while time.time() < deadline:
            run = client.get(f"/api/runs/{run_id}").json()
            if run["status"] in ("done", "failed"):
                break
            time.sleep(0.25)
        yield client, run_id, handler, tmp
    root.removeHandler(handler)
    root.setLevel(prev)


def _case_id(client, number):
    rows = client.get("/api/cases", params={"q": number}).json()["rows"]
    return rows[0]["id"]


def test_run_completes(env):
    client, run_id, _, _ = env
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["status"] == "done" and run["processed"] == run["total"] == 19 and run["failed"] == 0
    assert run["provider"] == "mock" and run["prompt_versions"]["troubleshooting"].startswith("troubleshooting.v1")


def test_review_routing_via_api(env):
    client, run_id, _, _ = env
    q = client.get("/api/review-queue", params={"run_id": run_id}).json()
    groups = {g["reason"]: {c["case_number"] for c in g["cases"]} for g in q["groups"]}
    assert {"00100002", "00100004", "00100005", "00100006", "00100015"} <= groups["RULE_BREACH"]  # SLO
    assert "00100011" in groups["RULE_BREACH"]                                                     # 3-strike
    assert "00100013" in groups["HOT_CUSTOMER"]
    all_queued = set().union(*groups.values())
    assert "00100001" not in all_queued
    detail = client.get(f"/api/cases/{_case_id(client, '00100011')}").json()
    codes = [r["code"] for r in detail["audit"]["review_reasons"]]
    assert "RULE_BREACH" in codes and any("3-strike" in r["detail"] for r in detail["audit"]["review_reasons"])


def test_queue_sorting(env):
    client, _, _, _ = env
    by_score = client.get("/api/review-queue", params={"sort": "score"}).json()
    rows = [c for g in by_score["groups"] for c in g["cases"]]
    assert rows
    by_sev = client.get("/api/review-queue", params={"sort": "severity"}).json()["groups"][0]["cases"]
    sevs = [c["severity"] or 9 for c in by_sev]
    assert sevs == sorted(sevs)


def test_case_filters(env):
    client, _, _, _ = env
    sev1 = client.get("/api/cases", params={"severity": 1}).json()["rows"]
    assert {r["case_number"] for r in sev1} == {"00100002", "00100003"}
    low = client.get("/api/cases", params={"score_max": 5}).json()["rows"]
    assert all(r["overall"] is not None and r["overall"] <= 5 for r in low)
    hot = client.get("/api/cases", params={"review_reason": "HOT_CUSTOMER"}).json()["rows"]
    assert "00100013" in {r["case_number"] for r in hot}


def test_reviewer_actions_never_overwrite_machine_result(env):
    client, _, _, _ = env
    detail = client.get(f"/api/cases/{_case_id(client, '00100013')}").json()
    audit_id, machine = detail["audit"]["id"], detail["audit"]["overall"]
    assert client.post(f"/api/audits/{audit_id}/review-actions",
                       json={"action": "override", "reviewer_name": "QA Lead"}).status_code == 422
    r = client.post(f"/api/audits/{audit_id}/review-actions",
                    json={"action": "override", "reviewer_name": "QA Lead", "score_override": 2.5,
                          "comment": "Repeated log requests confirmed."})
    assert r.status_code == 201
    client.post(f"/api/audits/{audit_id}/review-actions",
                json={"action": "approve", "reviewer_name": "QA Lead"})
    after = client.get(f"/api/cases/{_case_id(client, '00100013')}").json()
    assert after["audit"]["overall"] == machine
    trail = after["audit"]["review_actions"]
    assert [x["action"] for x in trail] == ["override", "approve"]
    assert after["review"]["score_override"] == 2.5


def test_dashboard_shape(env):
    client, _, _, _ = env
    d = client.get("/api/dashboard").json()
    assert d["kpis"]["cases"] == 19 and d["kpis"]["review_queue"] > 0
    assert sum(b["count"] for b in d["score_distribution"]) == 19
    assert {r["severity"] for r in d["slo_by_severity"]} == {1, 2, 3, 4}
    assert d["three_strike"]["APPLIED_INCORRECTLY"] == 1 and d["three_strike"]["APPLIED_CORRECTLY"] == 2


def test_upload_and_run(env, fixture_bundle):
    client, _, _, tmp = env
    data = fixture_bundle[0].read_bytes()
    r = client.post("/api/uploads", files={"file": ("export.xlsx", data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200
    st = r.json()["structure"]["mapping_status"]
    assert all(v == "found" for sheet in st.values() for v in sheet.values())
    assert client.post("/api/uploads", files={"file": ("x.csv", b"a,b", "text/csv")}).status_code == 400
    run = client.post("/api/runs", json={"source": "upload", "upload_id": r.json()["id"]}).json()
    deadline = time.time() + 120
    while time.time() < deadline and client.get(f"/api/runs/{run['id']}").json()["status"] not in ("done", "failed"):
        time.sleep(0.25)
    assert client.get(f"/api/runs/{run['id']}").json()["status"] == "done"


def test_no_fixture_pii_anywhere(env, fixture_bundle):
    client, run_id, handler, _ = env
    pii = [p.lower() for p in fixture_bundle[1]["pii"]]

    def assert_clean(blob: str, where: str):
        low = blob.lower()
        leaked = [p for p in pii if p in low]
        assert not leaked, f"{len(leaked)} planted value(s) found in {where}"

    # API responses
    bodies = [client.get(u).text for u in ("/api/dashboard", "/api/cases", "/api/review-queue", "/api/runs",
                                           "/api/config", f"/api/runs/{run_id}")]
    for row in client.get("/api/cases").json()["rows"]:
        bodies.append(client.get(f"/api/cases/{row['id']}").text)
    assert_clean("\n".join(bodies), "API responses")

    # Database: every text-ish column of every table
    dump = []
    with get_engine().connect() as conn:
        for table in inspect(conn).get_table_names():
            for rec in conn.execute(text(f"SELECT * FROM {table}")):
                dump.append(" ".join(str(v) for v in rec))
    assert_clean("\n".join(dump), "database")

    # Export
    xl = client.get("/api/export.xlsx")
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb.sheetnames == ["Summary", "Cases", "Findings", "Review Actions"]
    cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    assert_clean("\n".join(cells), "export")
    assert wb["Findings"].max_row > 1

    # Logs
    assert handler.lines
    assert_clean("\n".join(handler.lines), "logs")
