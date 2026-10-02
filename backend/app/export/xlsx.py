"""Excel export: Summary, Cases, Findings (with evidence), Review Actions. Redacted data only."""
from __future__ import annotations

import io
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db import tables as t

HEADER_FILL = PatternFill("solid", fgColor="1E1E24")
HEADER_FONT = Font(bold=True, color="F2F2F4")
DIMS = ["troubleshooting", "communication", "slo", "idle", "three_strike", "temperature_handling"]


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None][:200])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    return ws


def _naive(dt):
    return dt.replace(tzinfo=None) if dt else None


def build_export(run: t.Run, cases: list[t.Case]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    audits = [c.audit for c in cases if c.audit]
    scored = [a.overall for a in audits if a.overall is not None]
    reasons = Counter(r["code"] for a in audits for r in {x["code"]: x for x in a.review_reasons}.values())
    summary = [
        ("Run ID", run.id), ("Source", run.source), ("Synthetic data", run.synthetic),
        ("As of (UTC)", _naive(run.as_of)), ("Provider", run.provider), ("Model", run.model),
        ("Temperature", "model default" if run.temperature is None else run.temperature),
        ("Config hash", run.config_hash),
        ("Prompt versions", ", ".join(f"{k}={v}" for k, v in sorted((run.prompt_versions or {}).items()))),
        ("Cases audited", len(cases)),
        ("Average overall /10", round(sum(scored) / len(scored), 2) if scored else None),
        ("Cases in review queue", sum(1 for a in audits if a.needs_review)),
        ("EVAL_FAILED", sum(1 for a in audits if a.state == "EVAL_FAILED")),
        ("REDACTION_FAILED", sum(1 for a in audits if a.state == "REDACTION_FAILED")),
    ] + [(f"Review reason: {k}", v) for k, v in sorted(reasons.items())]
    _sheet(wb, "Summary", ["Metric", "Value"], [list(r) for r in summary])

    case_rows = []
    for c in cases:
        a = c.audit
        d = {x["dimension"]: x for x in (a.dimensions if a else [])}
        case_rows.append([
            c.case_number, c.severity, c.status, c.state, c.owner_label, c.account_label,
            _naive(c.opened_at), _naive(c.closed_at), a.overall if a else None,
            *[(d[k]["score"] if d.get(k) and d[k]["status"] == "SCORED" else "excluded") for k in DIMS],
            (a.slo or {}).get("status") if a else None, (a.slo or {}).get("actual_minutes") if a else None,
            (a.slo or {}).get("target_minutes") if a else None,
            (a.idle or {}).get("status") if a else None, (a.idle or {}).get("support_idle_hours") if a else None,
            (a.three_strike or {}).get("status") if a else None, (a.three_strike or {}).get("reason") if a else None,
            a.temperature_value if a else None, a.trajectory if a else None,
            a.confidence_level if a else None, a.confidence_score if a else None,
            "; ".join(f"{r['code']}: {r['detail']}" for r in (a.confidence_reasons if a else [])),
            ", ".join(sorted({r["code"] for r in (a.review_reasons if a else [])})),
            a.model if a else None,
        ])
    _sheet(wb, "Cases", [
        "Case", "Severity", "Status", "Audit state", "Engineer", "Account", "Opened (UTC)", "Closed (UTC)",
        "Overall /10", "Troubleshooting /10", "Communication /10", "SLO /10", "Idle /10", "3-strike /10",
        "Temperature handling /10", "SLO result", "SLO actual min", "SLO target min", "Idle result",
        "Support idle hours", "3-strike result", "3-strike reason", "Temperature 1-5", "Trajectory",
        "Confidence", "Confidence score", "Confidence reasons", "Review reasons", "Model",
    ], case_rows)

    finding_rows = []
    for c in cases:
        if not c.audit:
            continue
        items = {i.ref_id: i for i in c.items}
        for f in c.audit.findings:
            for ev in f.evidence or [{"ref_id": None, "timestamp": None}]:
                item = items.get(ev.get("ref_id"))
                excerpt = (item.body[:300] if item else (c.resolution[:300] if ev.get("ref_id") == "RES"
                           else c.description[:300] if ev.get("ref_id") == "DESC" else ""))
                finding_rows.append([c.case_number, f.dimension, f.kind, f.text, ev.get("ref_id"),
                                     ev.get("timestamp"), excerpt])
    _sheet(wb, "Findings", ["Case", "Dimension", "Kind", "Finding", "Evidence ref", "Evidence timestamp",
                            "Evidence excerpt (redacted)"], finding_rows)

    review_rows = [
        [c.case_number, c.audit.overall, x.action, x.score_override, x.reviewer_name, _naive(x.created_at), x.comment]
        for c in cases if c.audit for x in c.audit.review_actions
    ]
    _sheet(wb, "Review Actions", ["Case", "Machine overall /10", "Action", "Score override", "Reviewer",
                                  "At (UTC)", "Comment"], review_rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
